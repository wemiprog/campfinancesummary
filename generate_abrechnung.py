#!/usr/bin/env python3
"""HeadWaters Lagerabrechnung

Liest die Kontoauszüge einer Kostenstelle und baut eine PDF-Abrechnung:
die ersten Seiten als Übersicht, danach die Kontoauszug-PDFs dieser
Kostenstelle, am Schluss die Belege (nach Ausgabengrösse, Lebensmittel
zuletzt). Als Erstes fragt das Skript den Lagersteckbrief ab.

Start:
    uv run generate_abrechnung.py

Die Antworten landen in lagerinfo.json und gelten beim nächsten
Lauf als Vorschlag.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pymupdf


# ─────────────────────────────────────────────────────────────
# Pfade
# ─────────────────────────────────────────────────────────────

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
OUTPUT = ROOT / "output"
ASSETS = ROOT / "assets"
LOGO_DATEI = ASSETS / "headwaters-logo.png"
LAGERINFO_DATEI = ROOT / "lagerinfo.json"

MONATE = [
    "",
    "Januar",
    "Februar",
    "März",
    "April",
    "Mai",
    "Juni",
    "Juli",
    "August",
    "September",
    "Oktober",
    "November",
    "Dezember",
]


# ─────────────────────────────────────────────────────────────
# Kontenplan (wie in der Buchhaltung, mit Umlauten)
# ─────────────────────────────────────────────────────────────

KONTO_NAMEN = {
    "3200": "Shop-Verkäufe",
    "3220": "Rechnungen",
    "3300": "Mitgliederbeiträge",
    "3400": "Spenden",
    "3500": "Teilnehmerbeiträge",
    "3600": "Weiterleitungen",
    "3610": "Übrig vom Vorjahr",
    "3611": "Übrig am Jahresende",
    "3900": "Weitere Einnahmen",
    "4000": "Materialkauf",
    "4100": "Geräteeinkauf",
    "4120": "Medienkauf",
    "4130": "Transfer Karte TJ",
    "4140": "Rückzahlungen",
    "4150": "Persönliche Projektmittel",
    "4200": "Lebensmittel",
    "4210": "Verbrauchsmaterial",
    "4300": "Mietkosten",
    "4410": "Serverkosten",
    "4411": "DNS-Kosten",
    "4413": "Sonstige IT-Aufwände",
    "4420": "Porto & Onlinedruck",
    "4421": "Email, SMS",
    "4430": "Druckkosten",
    "4500": "Spesen",
    "4510": "Reisekosten",
    "4900": "Unterstützung",
    "6100": "Buchhaltungssoftware",
    "6200": "Bankgebühren",
    "6300": "Transaktionsgebühren",
    "6700": "Bewilligungen, Sonstiges",
    "6950": "Währungsdifferenzen",
}

AUFWAND_GRUPPEN = [
    ("4 Aufwand", ["4000", "4100", "4120", "4500", "4510"]),
    ("42–43 Essen & Locations", ["4200", "4210", "4300"]),
    (
        "44 Bezogene Dienstleistungen",
        ["4410", "4411", "4413", "4420", "4421", "4430"],
    ),
    ("46 Finanzaufwand", ["4900", "4130", "4140", "4150"]),
    (
        "6 Übriger Aufwand",
        ["6100", "6200", "6300", "6700", "6950"],
    ),
]

ERTRAG_GRUPPEN = [
    ("32 Dienstleistungen", ["3200", "3220"]),
    ("33–34 Spenden und Mitgliederbeiträge", ["3300", "3400"]),
    ("35 Teilnehmerbeiträge", ["3500"]),
    (
        "36 Weiterleitungen und Vorjahresbeträge",
        ["3600", "3610", "3611", "3900"],
    ),
]

# UBS eBill: 40 Rappen pro eingehende Zahlung
EBILL_GEBUEHR = 0.40


# ─────────────────────────────────────────────────────────────
# Datenstrukturen
# ─────────────────────────────────────────────────────────────

@dataclass
class Buchung:
    id: str
    datum: datetime | None
    text: str
    beleg_nr: str
    betrag: float
    soll: str
    haben: str
    kostenstelle: str


@dataclass
class Konto:
    nummer: str
    name: str
    kategorie: str  # Aufwand / Ertrag / Aktiv / Passiv
    gruppe: str  # Ordnername, z. B. "42-43 Essen & Locations"
    xlsx: Path | None = None
    pdf: Path | None = None
    soll: float = 0.0
    haben: float = 0.0
    n_buchungen: int = 0

    @property
    def ist_aufwand(self) -> bool:
        return self.nummer[:1] in "456"

    @property
    def ist_ertrag(self) -> bool:
        return self.nummer.startswith("3")

    @property
    def saldo(self) -> float:
        """Ertrag: Haben − Soll. Aufwand: Soll − Haben."""
        if self.ist_ertrag:
            return round(self.haben - self.soll, 2)
        if self.ist_aufwand:
            return round(self.soll - self.haben, 2)
        return round(self.soll - self.haben, 2)

    @property
    def hat_buchungen(self) -> bool:
        return self.n_buchungen > 0 and (self.soll != 0 or self.haben != 0)


@dataclass
class Lagerinfo:
    """Alles, was mer us de Kontoauszüg nid chöi läse."""

    name: str
    kostenstelle: str
    teilnehmer: int  # Erwachsene / Jugendliche, ohne Kinder
    kinder: int
    kind_faktor: float  # Kinder zählen als so viele Personen
    anreise: date
    abreise: date
    verpflegungstage: float
    ort: str

    @property
    def naechte(self) -> int:
        return (self.abreise - self.anreise).days

    @property
    def jahr(self) -> str:
        return str(self.anreise.year)

    @property
    def datum_text(self) -> str:
        return format_spanne(self.anreise, self.abreise)

    @property
    def personen(self) -> float:
        """Personen, Kinder anteilig gerechnet."""
        return self.teilnehmer + self.kinder * self.kind_faktor

    @property
    def koepfe(self) -> int:
        return self.teilnehmer + self.kinder

    def as_dict(self) -> dict:
        return {
            "name": self.name,
            "kostenstelle": self.kostenstelle,
            "teilnehmer": self.teilnehmer,
            "kinder": self.kinder,
            "kind_faktor": self.kind_faktor,
            "anreise": self.anreise.isoformat(),
            "abreise": self.abreise.isoformat(),
            "verpflegungstage": self.verpflegungstage,
            "ort": self.ort,
        }

    @classmethod
    def from_dict(cls, data: dict) -> "Lagerinfo":
        return cls(
            name=data["name"],
            kostenstelle=data["kostenstelle"],
            teilnehmer=int(data["teilnehmer"]),
            kinder=int(data["kinder"]),
            kind_faktor=float(data["kind_faktor"]),
            anreise=date.fromisoformat(data["anreise"]),
            abreise=date.fromisoformat(data["abreise"]),
            verpflegungstage=float(data["verpflegungstage"]),
            ort=data["ort"],
        )


@dataclass
class Zahlungsmethode:
    name: str
    volumen: float
    gebuehren: float
    anzahl: int = 0
    hinweis: str = ""
    satz: str = ""


@dataclass
class Auswertung:
    lager: Lagerinfo
    buchungen: list[Buchung]
    konten: dict[str, Konto]
    aufwand_total: float
    ertrag_total: float
    ertrag_ohne_vorjahr: float
    ertrag_ohne_vorjahr_ohne_spenden: float
    spenden: float
    vorjahr: float
    saldo_jahr: float
    ueberschuss_jahr: float
    uebriges_geld: float
    zahlungen: list[Zahlungsmethode]
    kosten_nach_konto: list[tuple[str, str, float]]


# ─────────────────────────────────────────────────────────────
# Kleine Helfer
# ─────────────────────────────────────────────────────────────

def chf(betrag: float, mit_waehrung: bool = True) -> str:
    """Schweizer Zahlenformat: 19'725.77 CHF"""
    vorzeichen = "−" if betrag < 0 else ""
    absolut = abs(round(betrag, 2))
    formatiert = f"{absolut:,.2f}".replace(",", "'")
    if mit_waehrung:
        return f"{vorzeichen}{formatiert} CHF"
    return f"{vorzeichen}{formatiert}"


def konto_nr(bezeichnung: str | None) -> str:
    if not bezeichnung:
        return ""
    return str(bezeichnung).split()[0]


def schoener_kontoname(roh: str) -> str:
    """Dateinamen wie 'Geraeteeinkauf' → 'Geräteeinkauf'."""
    ersatz = [
        ("ae", "ä"),
        ("oe", "ö"),
        ("ue", "ü"),
        ("Ae", "Ä"),
        ("Oe", "Ö"),
        ("Ue", "Ü"),
    ]
    out = roh
    for alt, neu in ersatz:
        out = out.replace(alt, neu)
    return out.replace("_", " ").strip()


def parse_konto_aus_dateiname(path: Path) -> tuple[str, str]:
    """Liefert (Kontonummer, Name) aus z. B. '4300 Mietkosten.xlsx'.

    Der Gruppen-Präfix kann selbst eine Zahl enthalten
    ('109 Transferkonten 1092 SumUp …') — darum die *letzte*
    3–4-stellige Zahl, das ist die Kontonummer.
    """
    stem = path.stem
    treffer = list(re.finditer(r"\d{3,4}", stem))
    if not treffer:
        return "", schoener_kontoname(stem)
    last = treffer[-1]
    nummer = last.group(0)
    name = stem[last.end() :].strip(" _-|")
    return nummer, KONTO_NAMEN.get(nummer, schoener_kontoname(name) or nummer)


def finde_ordner(stichwort: str) -> Path:
    treffer = [
        p
        for p in DATA.iterdir()
        if p.is_dir() and stichwort in p.name
    ]
    if not treffer:
        raise FileNotFoundError(
            f"Kei Ordner mit «{stichwort}» unter {DATA}"
        )
    return sorted(treffer, key=lambda p: p.name)[0]


def frage(prompt: str, default: str | None = None, pflicht: bool = False) -> str:
    suffix = f" [{default}]" if default not in (None, "") else ""
    while True:
        try:
            roh = input(f"{prompt}{suffix}: ").strip()
        except EOFError:
            roh = ""
        if not roh:
            if default not in (None, ""):
                return str(default)
            if pflicht:
                print("   → Bitte öppis iigeh.")
                continue
            return ""
        return roh


def frage_zahl(prompt: str, default: float | int | None = None, ganzzahl: bool = False):
    while True:
        default_txt = None if default is None else str(default)
        roh = frage(prompt, default_txt, pflicht=default is None)
        roh = roh.replace("'", "").replace(" ", "").replace(",", ".")
        try:
            wert = int(roh) if ganzzahl else float(roh)
        except ValueError:
            print("   → Das isch kä Zahl.")
            continue
        return wert


def parse_datum(text: str) -> date:
    """Nimmt 22.5.2026, 22.05.26, 2026-05-22, …"""
    roh = text.strip()
    iso = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})", roh)
    if iso:
        return date(int(iso[1]), int(iso[2]), int(iso[3]))
    m = re.fullmatch(r"(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{2}|\d{4})", roh)
    if not m:
        raise ValueError(roh)
    tag, monat, jahr = int(m[1]), int(m[2]), int(m[3])
    if jahr < 100:
        jahr += 2000
    return date(jahr, monat, tag)


def format_datum(d: date) -> str:
    return f"{d.day:02d}.{d.month:02d}.{d.year}"


def format_spanne(start: date, ende: date) -> str:
    if start.year == ende.year and start.month == ende.month:
        return f"{start.day}.–{ende.day}. {MONATE[start.month]} {start.year}"
    if start.year == ende.year:
        return (
            f"{start.day}. {MONATE[start.month]} – "
            f"{ende.day}. {MONATE[ende.month]} {start.year}"
        )
    return f"{format_datum(start)} – {format_datum(ende)}"


def frage_datum(prompt: str, default: date | None = None) -> date:
    default_txt = format_datum(default) if default else None
    while True:
        roh = frage(prompt, default_txt, pflicht=True)
        try:
            return parse_datum(roh)
        except ValueError:
            print("   → Bitte als TT.MM.JJJJ, z. B. 22.05.2026")


def lagerinfo_laden() -> Lagerinfo | None:
    if not LAGERINFO_DATEI.exists():
        return None
    try:
        data = json.loads(LAGERINFO_DATEI.read_text(encoding="utf-8"))
        return Lagerinfo.from_dict(data)
    except (json.JSONDecodeError, KeyError, ValueError) as exc:
        print(f"  Hinweis: {LAGERINFO_DATEI.name} isch nid lesbar ({exc}).")
        return None


def lagerinfo_speichern(lager: Lagerinfo) -> None:
    LAGERINFO_DATEI.write_text(
        json.dumps(lager.as_dict(), indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )


def lager_zusammenfassung(lager: Lagerinfo) -> str:
    kind_txt = "½" if abs(lager.kind_faktor - 0.5) < 0.001 else str(lager.kind_faktor)
    return (
        f"  {lager.name}\n"
        f"  {lager.ort}\n"
        f"  {lager.datum_text}\n"
        f"  {lager.naechte} Nächte · {lager.verpflegungstage:g} Verpflegungstage\n"
        f"  {lager.teilnehmer} Erwachsene + {lager.kinder} Kinder "
        f"= {lager.koepfe} Köpfe, {lager.personen:g} Personen (Kinder {kind_txt})"
    )


# ─────────────────────────────────────────────────────────────
# Excel lesen
# ─────────────────────────────────────────────────────────────

def lies_journal(xlsx_dir: Path, kostenstelle: str | None = None) -> list[Buchung]:
    pfad = xlsx_dir / "Buchungsjournal.xlsx"
    wb = openpyxl.load_workbook(pfad, data_only=True)
    ws = wb.active
    buchungen: list[Buchung] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        ks = (row[8] or "").strip() if row[8] else ""
        if kostenstelle and ks.upper() != kostenstelle.upper():
            continue
        buchungen.append(
            Buchung(
                id=str(row[0] or ""),
                datum=row[1] if isinstance(row[1], datetime) else None,
                text=str(row[2] or ""),
                beleg_nr=str(row[3] or ""),
                betrag=float(row[5] or 0),
                soll=str(row[6] or ""),
                haben=str(row[7] or ""),
                kostenstelle=ks,
            )
        )
    return buchungen


def lies_kontoauszug_xlsx(pfad: Path, kostenstelle: str) -> tuple[int, float, float]:
    """Anzahl echter Buchungen plus Soll/Haben-Summen."""
    wb = openpyxl.load_workbook(pfad, data_only=True)
    ws = wb.active
    n, soll, haben = 0, 0.0, 0.0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(row):
            continue
        datum = row[0]
        if not isinstance(datum, datetime):
            continue  # Anfangssaldo / Leerzeilen
        ks = ""
        if len(row) > 10 and row[10]:
            ks = str(row[10]).strip()
        if ks and ks.upper() != kostenstelle.upper():
            continue
        betrag_a = float(row[4] or 0)
        betrag_b = float(row[5] or 0)
        if betrag_a == 0 and betrag_b == 0:
            continue
        n += 1
        soll += betrag_a
        haben += betrag_b
    return n, round(soll, 2), round(haben, 2)


def sammle_konten(xlsx_dir: Path, pdf_dir: Path, kostenstelle: str) -> dict[str, Konto]:
    konten: dict[str, Konto] = {}

    for xlsx in sorted(xlsx_dir.rglob("*.xlsx")):
        if xlsx.name.startswith(".") or xlsx.name == "Buchungsjournal.xlsx":
            continue
        nummer, name = parse_konto_aus_dateiname(xlsx)
        if not nummer:
            continue
        rel = xlsx.relative_to(xlsx_dir)
        kategorie = rel.parts[0] if rel.parts else ""
        gruppe = xlsx.parent.name
        n, soll, haben = lies_kontoauszug_xlsx(xlsx, kostenstelle)
        konten[nummer] = Konto(
            nummer=nummer,
            name=KONTO_NAMEN.get(nummer, name),
            kategorie=kategorie,
            gruppe=gruppe,
            xlsx=xlsx,
            soll=soll,
            haben=haben,
            n_buchungen=n,
        )

    for pdf in sorted(pdf_dir.rglob("*.pdf")):
        nummer, name = parse_konto_aus_dateiname(pdf)
        if not nummer:
            continue
        if nummer in konten:
            konten[nummer].pdf = pdf
            if konten[nummer].name in (nummer, "") or "ae" in pdf.stem:
                konten[nummer].name = KONTO_NAMEN.get(nummer, name)
        else:
            rel = pdf.relative_to(pdf_dir)
            konten[nummer] = Konto(
                nummer=nummer,
                name=KONTO_NAMEN.get(nummer, name),
                kategorie=rel.parts[0] if rel.parts else "",
                gruppe=pdf.parent.name,
                pdf=pdf,
            )

    return konten


def pl_aus_journal(buchungen: list[Buchung]) -> dict[str, dict]:
    """Soll/Haben je Erfolgs-Konto aus dem Journal."""
    pl: dict[str, dict] = {}
    for b in buchungen:
        for seite, konto in (("soll", b.soll), ("haben", b.haben)):
            nr = konto_nr(konto)
            if not nr or nr[0] not in "3456":
                continue
            eintrag = pl.setdefault(
                nr, {"soll": 0.0, "haben": 0.0, "name": konto}
            )
            eintrag[seite] += b.betrag
            eintrag["name"] = konto
    return pl


# ─────────────────────────────────────────────────────────────
# Zahlungsmethoden
# ─────────────────────────────────────────────────────────────

BANK_KONTEN = {"1000", "1021", "1022", "1024", "1090", "1091"}
ANDERE_LAGER = ("fbc", "iglu", "ybc", "juko", "obc", "jidun", "diagonal", "wes")


def ist_ebill_betreff(text: str) -> bool:
    """eBill-Zahlungen erkenne mer am Keyword «ebill» im Buchungstext / Betreff."""
    t = (text or "").lower().replace("-", "").replace(" ", "")
    return "ebill" in t


def paycodes_aus_buchungen(buchungen: list[Buchung]) -> set[str]:
    """Vierstelliger alphanumerischer Paycode (Beleg-Nr. oder am Schluss vom Titel)."""
    codes: set[str] = set()
    for b in buchungen:
        beleg = (b.beleg_nr or "").strip().upper()
        if re.fullmatch(r"[A-Z0-9]{4}", beleg):
            codes.add(beleg)
        m = re.search(r"[-–]\s*([A-Z0-9]{4})\s*$", b.text.strip())
        if m:
            codes.add(m.group(1).upper())
    return codes


def rechnungsnummern_aus_xlsx(xlsx_dir: Path, kostenstelle: str) -> set[str]:
    """Rechnungsnummern us de Kontoauszüg (4+ Ziffere)."""
    nummern: set[str] = set()
    for f in xlsx_dir.rglob("*.xlsx"):
        if "3500" not in f.name and "1100" not in f.name:
            continue
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb.active
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        try:
            i_ks = headers.index("Kostenstelle")
            i_rid = headers.index("Rechnungs ID")
        except ValueError:
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            ks = str(row[i_ks] or "").strip().upper()
            if ks != kostenstelle.upper():
                continue
            rid = str(row[i_rid] or "").strip()
            if rid.isdigit() and len(rid) >= 4:
                nummern.add(rid)
    return nummern


def text_hat_paycode(text: str, codes: set[str]) -> bool:
    compact = re.sub(r"\s+", "", text.upper())
    for code in codes:
        if re.search(rf"(?<![A-Z0-9]){re.escape(code)}(?![A-Z0-9])", compact):
            return True
    return False


def text_hat_rechnungsnr(text: str, nummern: set[str]) -> bool:
    for nr in nummern:
        if re.search(rf"(?<!\d){re.escape(nr)}(?!\d)", text):
            return True
    return False


def ist_lager_hinweis(text: str) -> bool:
    """TWINT/RaiseNow het oft kei Paycode, aber «BTP» / «Pfingsten» im Text."""
    t = text.lower()
    if "spende" in t:
        return False
    if any(x in t for x in ANDERE_LAGER) and not re.search(r"\bbtp\b", t):
        return False
    return bool(re.search(r"\bbtp\b", t) or "pfingst" in t or "bibeltage" in t)


def gehoert_zum_lager(
    b: Buchung, paycodes: set[str], rechnungen: set[str]
) -> bool:
    if text_hat_paycode(b.text, paycodes) or text_hat_rechnungsnr(b.text, rechnungen):
        return True
    soll = konto_nr(b.soll)
    # RaiseNow/TWINT: Name + «BTP», ohni Paycode
    if soll == "1091" and ist_lager_hinweis(b.text):
        return True
    # UBS/Kasse ohni Paycode, aber klar BTP
    if soll in {"1021", "1000"} and ist_lager_hinweis(b.text):
        return True
    return False


def analysiere_zahlungen(
    buchungen: list[Buchung],
    eingaenge: list[Buchung],
    rechnungen: set[str] | None = None,
) -> list[Zahlungsmethode]:
    """Wie d Teilnehmerbeiträg aacho si — Überwysig mit/ohni eBill separat."""
    paycodes = paycodes_aus_buchungen(buchungen)
    rechnungen = rechnungen or set()

    sumup_volumen = 0.0
    sumup_anzahl = 0
    sumup_gebuehren = 0.0
    twint_gebuehren = 0.0
    twint_satz = 0.025
    twint_volumen = 0.0
    twint_anzahl = 0
    ebill_volumen = 0.0
    ebill_anzahl = 0
    ohne_ebill = 0.0
    ohne_ebill_anzahl = 0

    for b in buchungen:
        text = b.text.lower()
        soll_nr = konto_nr(b.soll)
        haben_nr = konto_nr(b.haben)

        if soll_nr == "6300" and haben_nr == "1092":
            sumup_gebuehren += b.betrag
        elif soll_nr == "6300" and "twint" in text:
            twint_gebuehren += b.betrag
            m = re.search(r"(\d+[.,]\d+)\s*%", b.text)
            if m:
                twint_satz = float(m.group(1).replace(",", ".")) / 100

        # SumUp-Eingänge, die bereits eine Kostenstelle haben
        if soll_nr == "1092" and haben_nr == "1100":
            sumup_volumen += b.betrag
            sumup_anzahl += 1

    for b in eingaenge:
        soll_nr = konto_nr(b.soll)
        haben_nr = konto_nr(b.haben)
        # nume Iigäng, kei Usgäng/Rückerstattige
        if haben_nr in BANK_KONTEN:
            continue
        if not gehoert_zum_lager(b, paycodes, rechnungen):
            continue
        if soll_nr == "1091":
            twint_volumen += b.betrag
            twint_anzahl += 1
        elif soll_nr == "1092":
            # SumUp ohne Kostenstelle — zählt extra zu den KS-Buchungen
            sumup_volumen += b.betrag
            sumup_anzahl += 1
        elif soll_nr in {"1021", "1000", "1022"}:
            if ist_ebill_betreff(b.text):
                ebill_anzahl += 1
                ebill_volumen += b.betrag
            else:
                ohne_ebill_anzahl += 1
                ohne_ebill += b.betrag

    if twint_volumen == 0 and twint_gebuehren and twint_satz:
        twint_volumen = round(twint_gebuehren / twint_satz, 2)

    twint_volumen = round(twint_volumen, 2)
    sumup_volumen = round(sumup_volumen, 2)
    ebill_volumen = round(ebill_volumen, 2)
    ohne_ebill = round(ohne_ebill, 2)
    ebill_gebuehren = round(ebill_anzahl * EBILL_GEBUEHR, 2)

    methoden = [
        Zahlungsmethode(
            "TWINT",
            twint_volumen,
            twint_gebuehren,
            twint_anzahl,
            (
                f"{twint_anzahl} Zahlungen über RaiseNow/TWINT"
                if twint_anzahl
                else "Volumen aus Gebühr zurückgerechnet"
            )
            + (f" · {twint_satz * 100:.1f} %" if twint_gebuehren else ""),
            f"{twint_satz * 100:.2f} %",
        ),
        Zahlungsmethode(
            "SumUp (Karte)",
            sumup_volumen,
            sumup_gebuehren,
            sumup_anzahl,
            f"{sumup_anzahl} Kartenzahlungen über das SumUp-Transferkonto",
            f"{(sumup_gebuehren / sumup_volumen * 100) if sumup_volumen else 0:.2f} %",
        ),
        Zahlungsmethode(
            "Überweisung mit eBill",
            ebill_volumen,
            ebill_gebuehren,
            ebill_anzahl,
            f"{ebill_anzahl} Zahlungen mit «ebill» im Betreff à {chf(EBILL_GEBUEHR)}",
            f"{chf(EBILL_GEBUEHR)} / Zlg.",
        ),
        Zahlungsmethode(
            "Überweisung ohne eBill",
            ohne_ebill,
            0.0,
            ohne_ebill_anzahl,
            f"{ohne_ebill_anzahl} Zahlungen UBS / QR / Bar, ohne «ebill» im Betreff",
            "—",
        ),
    ]
    return [m for m in methoden if m.volumen or m.gebuehren]


# ─────────────────────────────────────────────────────────────
# Auswertung zusammenbauen
# ─────────────────────────────────────────────────────────────

def auswerten(
    lager: Lagerinfo,
    buchungen: list[Buchung],
    konten: dict[str, Konto],
    eingaenge: list[Buchung] | None = None,
    rechnungen: set[str] | None = None,
) -> Auswertung:
    pl = pl_aus_journal(buchungen)

    def netto(nr: str) -> float:
        if nr not in pl:
            return 0.0
        s, h = pl[nr]["soll"], pl[nr]["haben"]
        if nr.startswith("3"):
            return round(h - s, 2)
        return round(s - h, 2)

    # Kontensalden aus dem Journal übernehmen (präziser als die Auszüge)
    for nr, werte in pl.items():
        if nr not in konten:
            konten[nr] = Konto(
                nummer=nr,
                name=KONTO_NAMEN.get(nr, werte["name"]),
                kategorie="Ertrag" if nr.startswith("3") else "Aufwand",
                gruppe="",
            )
        if nr.startswith("3"):
            konten[nr].haben = werte["haben"]
            konten[nr].soll = werte["soll"]
        else:
            konten[nr].soll = werte["soll"]
            konten[nr].haben = werte["haben"]

    aufwand = round(sum(netto(nr) for nr in pl if not nr.startswith("3")), 2)
    ertrag = round(sum(netto(nr) for nr in pl if nr.startswith("3")), 2)
    vorjahr = netto("3610")
    spenden = netto("3400")
    ertrag_ohne_vorjahr = round(ertrag - vorjahr, 2)
    ertrag_ohne_spenden = round(ertrag_ohne_vorjahr - spenden, 2)

    # Saldo ohne Vorjahr und ohne Spenden — Spenden werden separat gezeigt
    saldo_jahr = round(ertrag_ohne_spenden - aufwand, 2)
    ueberschuss_jahr = round(saldo_jahr + spenden, 2)
    uebriges_geld = round(ueberschuss_jahr + vorjahr, 2)

    kosten = []
    for nr in sorted(pl):
        if nr.startswith("3"):
            continue
        betrag = netto(nr)
        if betrag:
            kosten.append((nr, KONTO_NAMEN.get(nr, pl[nr]["name"]), betrag))
    kosten.sort(key=lambda x: -x[2])

    return Auswertung(
        lager=lager,
        buchungen=buchungen,
        konten=konten,
        aufwand_total=aufwand,
        ertrag_total=ertrag,
        ertrag_ohne_vorjahr=ertrag_ohne_vorjahr,
        ertrag_ohne_vorjahr_ohne_spenden=ertrag_ohne_spenden,
        spenden=spenden,
        vorjahr=vorjahr,
        saldo_jahr=saldo_jahr,
        ueberschuss_jahr=ueberschuss_jahr,
        uebriges_geld=uebriges_geld,
        zahlungen=analysiere_zahlungen(buchungen, eingaenge or [], rechnungen),
        kosten_nach_konto=kosten,
    )


# ─────────────────────────────────────────────────────────────
# PDF-Seiten: leere weglassen, Rest als Bilder
# ─────────────────────────────────────────────────────────────

def seite_hat_buchung(text: str) -> bool:
    """True, wenn auf der Seite mindestens ein Buchungsdatum steht."""
    return bool(re.search(r"\d{2}\.\d{2}\.\d{4}", text))


def pdfs_zum_anhaengen(konten: dict[str, Konto]) -> list[Konto]:
    """Nur Konten mit echten Buchungen, schön sortiert."""

    def sort_key(k: Konto) -> tuple:
        kat = {"Aufwand": 0, "Ertrag": 1, "Aktiv": 2, "Passiv": 3}.get(k.kategorie, 9)
        try:
            nr = int(k.nummer)
        except ValueError:
            nr = 9999
        return (kat, nr)

    return [
        k
        for k in sorted(konten.values(), key=sort_key)
        if k.pdf and k.hat_buchungen
    ]


A4 = pymupdf.paper_rect("a4")
BELEGE_DIR = DATA / "Belege"
BILD_ENDUNGEN = {".jpg", ".jpeg", ".png", ".webp", ".tif", ".tiff"}

INK = (0.11, 0.11, 0.11)
MUTED = (0.42, 0.45, 0.50)
BLUE = (0.239, 0.431, 0.659)
GREEN = (0.239, 0.478, 0.290)
LOSS = (0.769, 0.361, 0.243)
PAPER = (0.953, 0.961, 0.969)
BAR_BG = (0.882, 0.902, 0.925)
FONT_REG = Path("/System/Library/Fonts/Supplemental/Arial.ttf")
FONT_BOLD = Path("/System/Library/Fonts/Supplemental/Arial Bold.ttf")


class Stift:
    """Zeichnet Text und Formen auf eine A4-Seite, mit Unicode-Schriften."""

    def __init__(self, page: pymupdf.Page):
        self.page = page
        self.reg = pymupdf.Font(fontfile=str(FONT_REG))
        self.bld = pymupdf.Font(fontfile=str(FONT_BOLD))
        page.insert_font("reg", fontfile=str(FONT_REG))
        page.insert_font("bld", fontfile=str(FONT_BOLD))
        self.ml = 44.0
        self.mr = page.rect.width - 44.0

    @property
    def cx(self) -> float:
        return (self.ml + self.mr) / 2

    def breite(self, text: str, size: float, bold: bool = False) -> float:
        font = self.bld if bold else self.reg
        return font.text_length(text, fontsize=size)

    def text(
        self,
        x: float,
        y: float,
        s: str,
        *,
        size: float = 10,
        bold: bool = False,
        color: tuple[float, float, float] = INK,
        align: str = "left",
    ) -> float:
        if not s:
            return 0.0
        w = self.breite(s, size, bold)
        if align == "right":
            x = x - w
        elif align == "center":
            x = x - w / 2
        self.page.insert_text(
            (x, y),
            s,
            fontname="bld" if bold else "reg",
            fontsize=size,
            color=color,
        )
        return w

    def kasten(
        self,
        rect: pymupdf.Rect,
        fill: tuple[float, float, float] = PAPER,
        radius: float | None = 0.04,
    ) -> None:
        if radius:
            self.page.draw_rect(rect, color=None, fill=fill, radius=radius)
        else:
            self.page.draw_rect(rect, color=None, fill=fill)

    def linie(
        self,
        x0: float,
        y: float,
        x1: float,
        color: tuple[float, float, float] = GREEN,
        width: float = 1.2,
    ) -> None:
        self.page.draw_line(
            pymupdf.Point(x0, y), pymupdf.Point(x1, y), color=color, width=width
        )

    def umbrechen(self, text: str, size: float, max_w: float, bold: bool = False) -> list[str]:
        woerter = text.split()
        zeilen: list[str] = []
        aktuell = ""
        for wort in woerter:
            probe = f"{aktuell} {wort}".strip()
            if self.breite(probe, size, bold) <= max_w:
                aktuell = probe
            else:
                if aktuell:
                    zeilen.append(aktuell)
                aktuell = wort
        if aktuell:
            zeilen.append(aktuell)
        return zeilen or [""]


def buchungs_id(b: Buchung) -> str:
    s = str(b.id).strip()
    return str(int(s)) if s.isdigit() else s


def beleg_id_aus_name(name: str) -> str | None:
    m = re.match(r"Beleg\s+0*(\d+)\b", name, re.I)
    return m.group(1) if m else None


def index_belege(ordner: Path) -> dict[str, Path]:
    """Buchungs-ID → eine Datei (PDF vor Bild, Unicode-Duplikate zusammen)."""
    roh: dict[str, list[Path]] = {}
    if not ordner.is_dir():
        return {}
    for p in ordner.iterdir():
        if not p.is_file() or p.name.startswith("."):
            continue
        bid = beleg_id_aus_name(p.name)
        if bid:
            roh.setdefault(bid, []).append(p)
    chosen: dict[str, Path] = {}
    for bid, files in roh.items():
        files = sorted(
            files,
            key=lambda p: (
                0 if p.suffix.lower() == ".pdf" else 1,
                len(unicodedata.normalize("NFC", p.name)),
                unicodedata.normalize("NFC", p.name),
            ),
        )
        chosen[bid] = files[0]
    return chosen


def ist_aufwandskonto(nr: str) -> bool:
    return bool(nr) and nr[:1] in "46"


def belege_zum_anhaengen(
    buchungen: list[Buchung], index: dict[str, Path]
) -> tuple[list[tuple[Buchung, Path]], list[tuple[Buchung, Path]]]:
    """Aufwandsbelege dieser Kostenstelle: zuerst nach Betrag, Lebensmittel zuletzt."""
    andere: list[tuple[Buchung, Path]] = []
    lebensmittel: list[tuple[Buchung, Path]] = []
    for b in buchungen:
        pfad = index.get(buchungs_id(b))
        if not pfad:
            continue
        soll, haben = konto_nr(b.soll), konto_nr(b.haben)
        if not (ist_aufwandskonto(soll) or ist_aufwandskonto(haben)):
            continue
        if soll == "4200" or haben == "4200":
            lebensmittel.append((b, pfad))
        else:
            andere.append((b, pfad))
    andere.sort(key=lambda x: -x[0].betrag)
    lebensmittel.sort(key=lambda x: -x[0].betrag)
    return andere, lebensmittel


def haenge_kontoauszuege_an(doc: pymupdf.Document, konten: list[Konto]) -> int:
    """Bestehende Kontoauszug-PDFs joinen, leere Seiten weglassen."""
    n = 0
    for konto in konten:
        if not konto.pdf:
            continue
        src = pymupdf.open(konto.pdf)
        keep = [i for i, page in enumerate(src) if seite_hat_buchung(page.get_text())]
        if not keep:
            src.close()
            continue
        src.select(keep)
        doc.insert_pdf(src)
        n += src.page_count
        src.close()
    return n


def _beleg_beschriftung(b: Buchung) -> str:
    nr = konto_nr(b.soll) or konto_nr(b.haben)
    name = KONTO_NAMEN.get(nr, b.soll or b.haben)
    return f"Beleg {buchungs_id(b)} · {nr} {name} · {chf(b.betrag)}"


def haenge_datei_an(doc: pymupdf.Document, pfad: Path, beschriftung: str) -> int:
    suf = pfad.suffix.lower()
    if suf == ".pdf":
        src = pymupdf.open(pfad)
        try:
            before = doc.page_count
            doc.insert_pdf(src)
            return doc.page_count - before
        finally:
            src.close()
    if suf in BILD_ENDUNGEN:
        page = doc.new_page(width=A4.width, height=A4.height)
        stift = Stift(page)
        stift.text(stift.ml, 36, beschriftung, size=9, color=MUTED)
        page.insert_image(
            pymupdf.Rect(36, 54, A4.width - 36, A4.height - 36),
            filename=str(pfad),
            keep_proportion=True,
        )
        return 1
    print(f"  überspringe Beleg {pfad.name}")
    return 0


def trennerseite(doc: pymupdf.Document, titel: str, untertitel: str = "") -> None:
    page = doc.new_page(width=A4.width, height=A4.height)
    stift = Stift(page)
    stift.kasten(page.rect, fill=PAPER, radius=0)
    stift.text(stift.cx, 390, titel, size=26, bold=True, align="center")
    if untertitel:
        stift.text(stift.cx, 418, untertitel, size=11, color=MUTED, align="center")


# ─────────────────────────────────────────────────────────────
# PDF-Übersicht zeichnen
# ─────────────────────────────────────────────────────────────

def _neue_seite(doc: pymupdf.Document) -> Stift:
    page = doc.new_page(width=A4.width, height=A4.height)
    return Stift(page)


def fusszeile(stift: Stift, lager: Lagerinfo, seite: int, total: int = 4) -> None:
    y = stift.page.rect.height - 28
    stift.text(stift.ml, y, f"HeadWaters · {lager.name}", size=8, color=MUTED)
    stift.text(stift.mr, y, f"{seite} / {total}", size=8, color=MUTED, align="right")


def zeichne_marke(stift: Stift, x_right: float, y_top: float) -> None:
    logo_h = 32.0
    logo_w = logo_h * 1772 / 1100
    img = pymupdf.Rect(x_right - logo_w - 86, y_top, x_right - 86, y_top + logo_h)
    stift.page.insert_image(img, filename=str(LOGO_DATEI), keep_proportion=True)
    tx = img.x1 + 7
    stift.text(tx, y_top + 13, "HEADWATERS", size=10.5, bold=True)
    stift.text(tx, y_top + 25, "QUELLGEBIET", size=6.2, color=MUTED)


def kopf(stift: Stift, lager: Lagerinfo, titel: str, untertitel: str | None = None) -> float:
    zeichne_marke(stift, stift.mr, 30)
    stift.text(stift.ml, 50, titel, size=26, bold=True)
    if untertitel is not None:
        stift.text(stift.ml, 72, untertitel, size=14, bold=True, color=(0.28, 0.28, 0.28))
        return 86
    return 64


def _zeilenfarbe(betrag: float, voll: tuple[float, float, float]) -> tuple[float, float, float]:
    if abs(betrag) < 0.005:
        return (0.62, 0.65, 0.68)
    return voll


def zeichne_pl_karte(
    stift: Stift,
    y0: float,
    kopftext: str,
    gruppen: list[tuple[str, list[str]]],
    konten: dict[str, Konto],
    total: float,
    total_label: str,
) -> float:
    pad = 16.0
    x0, x1 = stift.ml, stift.mr
    inner_l = x0 + pad
    inner_r = x1 - pad
    y = y0 + 22
    start_y = y0

    zeilen_h = 0.0
    for _, nummern in gruppen:
        zeilen_h += 16 + len(nummern) * 13 + 18
    hoehe = 28 + zeilen_h + 26
    stift.kasten(pymupdf.Rect(x0, y0, x1, y0 + hoehe), fill=PAPER, radius=0.025)

    stift.text(stift.cx, y, kopftext.upper(), size=9, color=MUTED, align="center")
    y += 18

    for gruppenname, nummern in gruppen:
        stift.text(inner_l, y, gruppenname, size=10.5, bold=True)
        y += 16
        gruppe_total = 0.0
        for nr in nummern:
            betrag = konten[nr].saldo if nr in konten else 0.0
            gruppe_total += betrag
            name = f"{nr} {KONTO_NAMEN.get(nr, konten[nr].name if nr in konten else nr)}"
            stift.text(
                inner_l + 12,
                y,
                name,
                size=9.5,
                color=_zeilenfarbe(betrag, BLUE),
            )
            stift.text(
                inner_r,
                y,
                chf(betrag),
                size=9.5,
                color=_zeilenfarbe(betrag, GREEN),
                align="right",
            )
            y += 13
        stift.linie(inner_l + 12, y - 9, inner_r, GREEN, 1.15)
        stift.text(
            inner_r,
            y + 2,
            chf(gruppe_total),
            size=10,
            bold=True,
            color=GREEN,
            align="right",
        )
        y += 18

    y += 4
    stift.text(inner_l, y, total_label, size=12, bold=True)
    stift.text(inner_r, y, chf(total), size=12, bold=True, color=GREEN, align="right")
    return start_y + hoehe


def zeichne_seite_aufwand(doc: pymupdf.Document, a: Auswertung) -> None:
    stift = _neue_seite(doc)
    lager = a.lager
    y = kopf(stift, lager, "Abrechnung", lager.name)
    meta = (
        f"Kostenstelle {lager.kostenstelle}  ·  {lager.ort}  ·  {lager.datum_text}  ·  "
        f"{lager.naechte} Nächte  ·  {lager.verpflegungstage:g} Verpflegungstage  ·  "
        f"{lager.koepfe} Personen ({lager.teilnehmer} Erwachsene, {lager.kinder} Kinder)"
    )
    for zeile in stift.umbrechen(meta, 8.5, stift.mr - stift.ml):
        stift.text(stift.ml, y, zeile, size=8.5, color=MUTED)
        y += 12
    y += 8
    zeichne_pl_karte(
        stift, y, "Aufwand", AUFWAND_GRUPPEN, a.konten, a.aufwand_total, "Total Aufwand"
    )
    fusszeile(stift, lager, 1)


def _fazit_zeile(
    stift: Stift,
    y: float,
    label: str,
    wert: str,
    *,
    negativ: bool,
    note: str | None = None,
) -> float:
    w_label = stift.breite(label + "  ", 13, True)
    w_wert = stift.breite(wert, 14.5, True)
    x = stift.cx - (w_label + w_wert) / 2
    stift.text(x, y, label + "  ", size=13, bold=True)
    stift.text(x + w_label, y, wert, size=14.5, bold=True, color=LOSS if negativ else GREEN)
    y += 14
    if note:
        for zeile in stift.umbrechen(note, 9, 420):
            stift.text(stift.cx, y + 4, zeile, size=9, color=MUTED, align="center")
            y += 12
        y += 6
    else:
        y += 10
    return y


def zeichne_seite_ertrag(doc: pymupdf.Document, a: Auswertung) -> None:
    stift = _neue_seite(doc)
    lager = a.lager
    y = zeichne_pl_karte(
        stift, 36, "Ertrag", ERTRAG_GRUPPEN, a.konten, a.ertrag_total, "Total Ertrag"
    )
    y += 28
    # kleine Welle
    cx, wy = stift.cx, y
    shape = stift.page.new_shape()
    shape.draw_bezier(
        pymupdf.Point(cx - 90, wy),
        pymupdf.Point(cx - 50, wy - 11),
        pymupdf.Point(cx - 10, wy + 11),
        pymupdf.Point(cx + 30, wy),
    )
    shape.draw_bezier(
        pymupdf.Point(cx + 30, wy),
        pymupdf.Point(cx + 55, wy - 10),
        pymupdf.Point(cx + 75, wy + 8),
        pymupdf.Point(cx + 90, wy - 2),
    )
    shape.finish(color=INK, width=1.6, closePath=False)
    shape.commit()
    y += 28
    stift.text(stift.cx, y, "Fazit", size=22, bold=True, align="center")
    y += 28
    y = _fazit_zeile(
        stift,
        y,
        "Saldo in diesem Jahr:",
        f"CHF {chf(a.saldo_jahr, False)}",
        negativ=a.saldo_jahr < 0,
        note=(
            f"ohne Vorjahr, ohne Lagerspenden · Betrieb {chf(a.ertrag_ohne_vorjahr_ohne_spenden)}"
            f" − Aufwand {chf(a.aufwand_total)}"
        ),
    )
    stift.text(
        stift.cx,
        y,
        f"Verwendete (allgemeine) Lagerspenden: CHF {chf(a.spenden, False)}",
        size=10,
        color=MUTED,
        align="center",
    )
    y += 22
    y = _fazit_zeile(
        stift,
        y,
        "Überschuss von diesem Jahr:",
        f"CHF {chf(a.ueberschuss_jahr, False)}",
        negativ=a.ueberschuss_jahr < 0,
    )
    vor_col = LOSS if a.vorjahr < 0 else GREEN
    prefix = "Überschuss vom Vorjahr: "
    wert = f"CHF {chf(a.vorjahr, False)}"
    w1 = stift.breite(prefix, 10)
    w2 = stift.breite(wert, 10, True)
    x = stift.cx - (w1 + w2) / 2
    stift.text(x, y, prefix, size=10, color=MUTED)
    stift.text(x + w1, y, wert, size=10, bold=True, color=vor_col)
    y += 24
    _fazit_zeile(
        stift,
        y,
        "Übriges Geld über alle Jahre:",
        f"CHF {chf(a.uebriges_geld, False)}",
        negativ=a.uebriges_geld < 0,
    )
    fusszeile(stift, lager, 2)


def zeichne_seite_pro_person(doc: pymupdf.Document, a: Auswertung) -> None:
    stift = _neue_seite(doc)
    lager = a.lager
    lebensmittel = a.konten["4200"].saldo if "4200" in a.konten else 0.0
    miete = a.konten["4300"].saldo if "4300" in a.konten else 0.0
    personen = lager.personen or 1
    tage = lager.verpflegungstage or 1
    naechte = lager.naechte or 1
    pro_tag = lebensmittel / personen / tage
    pro_nacht = miete / personen / naechte
    kind_txt = (
        "½" if abs(lager.kind_faktor - 0.5) < 0.001 else str(lager.kind_faktor).replace(".", ",")
    )

    def block(y: float, titel: str, zeilen: list[str], ergebnis: str, hint: str) -> float:
        stift.text(stift.cx, y, titel, size=20, bold=True, align="center")
        y += 22
        for z in zeilen:
            stift.text(stift.cx, y, z, size=12, align="center")
            y += 16
        stift.text(stift.cx, y + 4, "ergibt", size=12, bold=True, align="center")
        y += 22
        stift.text(stift.cx, y, ergebnis, size=14.5, bold=True, align="center")
        y += 18
        for z in stift.umbrechen(hint, 9.5, 430):
            stift.text(stift.cx, y, z, size=9.5, color=MUTED, align="center")
            y += 13
        return y + 18

    y = 70
    y = block(
        y,
        "Essenskosten",
        [
            f"Ausgaben für Lebensmittel: {chf(lebensmittel)}",
            f"Personen (Kinder {kind_txt} gerechnet): {personen:g}",
            f"Tage: {lager.verpflegungstage:g}  ({lager.datum_text})",
        ],
        f"{chf(pro_tag)} pro Person und Tag",
        "Die Küche wurde vom Haus übernommen — die Verpflegung steckt grösstenteils "
        "in den Mietkosten. Hier nur die extra Lebensmittel.",
    )
    y = block(
        y,
        "Übernachtungskosten",
        [
            f"Mietkosten: {chf(miete)}",
            f"Personen (Kinder {kind_txt} gerechnet): {personen:g}",
            f"Nächte: {lager.naechte}  ({lager.datum_text})",
        ],
        f"{chf(pro_nacht)} pro Person und Nacht",
        f"Teilzeitbesucher werden normal gerechnet. {lager.ort}",
    )
    stift.text(stift.cx, y, "Einzelheiten", size=20, bold=True, align="center")
    y += 22
    for z in stift.umbrechen(
        "Auf den nächsten Seiten sind die kompletten Kontoauszüge mit den konkreten "
        "Ausgaben und Einnahmen zu sehen.",
        11.5,
        430,
    ):
        stift.text(stift.cx, y, z, size=11.5, align="center")
        y += 16
    y += 6
    for z in stift.umbrechen(
        "Die Belege der Aufwände folgen am Schluss, sortiert nach Ausgabengrösse — "
        "die Lebensmittelbelege ganz am Ende.",
        11.5,
        430,
    ):
        stift.text(stift.cx, y, z, size=11.5, align="center")
        y += 16
    fusszeile(stift, lager, 3)


def _balken(
    stift: Stift,
    x: float,
    y: float,
    breite: float,
    anteil: float,
    color: tuple[float, float, float] = GREEN,
) -> None:
    h = 7.0
    stift.page.draw_rect(
        pymupdf.Rect(x, y, x + breite, y + h),
        color=None,
        fill=BAR_BG,
        radius=0.5,
    )
    fill_w = breite * max(0.0, min(1.0, anteil))
    if fill_w > 1:
        stift.page.draw_rect(
            pymupdf.Rect(x, y, x + fill_w, y + h),
            color=None,
            fill=color,
            radius=0.5,
        )


def zeichne_seite_zahlungen(doc: pymupdf.Document, a: Auswertung) -> None:
    stift = _neue_seite(doc)
    lager = a.lager
    kopf(stift, lager, "Zahlungen & Kosten")
    y_top = 78
    gap = 12
    mid = (stift.ml + stift.mr) / 2
    left0, left1 = stift.ml, mid - gap / 2
    right0, right1 = mid + gap / 2, stift.mr
    n_m = max(1, len(a.zahlungen))
    n_k = max(1, len(a.kosten_nach_konto))
    card_h = max(56 + 44 + n_m * 54 + 36, 56 + 40 + n_k * 40) + 12
    bottom = y_top + card_h
    stift.kasten(pymupdf.Rect(left0, y_top, left1, bottom), radius=0.03)
    stift.kasten(pymupdf.Rect(right0, y_top, right1, bottom), radius=0.03)

    pad = 12
    # linke Spalte
    x = left0 + pad
    xr = left1 - pad
    y = y_top + 22
    stift.text(x, y, "Zahlungsmethoden", size=12, bold=True)
    y += 16
    hint = (
        f"So sind die Teilnehmerbeiträge (netto {chf(sum(m.volumen for m in a.zahlungen))}) "
        f"angekommen. eBill kostet {chf(EBILL_GEBUEHR)} pro eingehende Überweisung."
    )
    for zeile in stift.umbrechen(hint, 8, xr - x):
        stift.text(x, y, zeile, size=8, color=MUTED)
        y += 11
    y += 8
    tb_vol = sum(m.volumen for m in a.zahlungen) or 1
    for m in a.zahlungen:
        label = f"{m.name}  ·  {m.anzahl}×" if m.anzahl else m.name
        stift.text(x, y, label, size=9.5, bold=True)
        stift.text(xr, y, chf(m.volumen), size=9.5, align="right")
        y += 5
        _balken(stift, x, y, xr - x, m.volumen / tb_vol, BLUE)
        y += 14
        sub = m.hinweis
        if m.gebuehren:
            sub += f" · Gebühren {chf(m.gebuehren)}"
        else:
            sub += " · keine Gebühr in der Kostenstelle"
        for zeile in stift.umbrechen(sub, 7.5, xr - x):
            stift.text(x, y, zeile, size=7.5, color=MUTED)
            y += 10
        y += 6

    y += 8
    stift.linie(x, y, xr, (0.84, 0.86, 0.89), 0.6)
    y += 14
    gebuehren_total = sum(m.gebuehren for m in a.zahlungen)
    stift.text(x, y, "Total Gebühren", size=9, bold=True)
    stift.text(xr, y, chf(gebuehren_total), size=9, bold=True, align="right")

    # rechte Spalte
    x = right0 + pad
    xr = right1 - pad
    y = y_top + 22
    stift.text(x, y, "Wohin das Geld ging", size=12, bold=True)
    y += 16
    personen = lager.personen or 1
    pro_kopf = a.aufwand_total / personen
    hint = f"Total Aufwand {chf(a.aufwand_total)} · {chf(pro_kopf)} pro Person (Kinder anteilig)."
    for zeile in stift.umbrechen(hint, 8, xr - x):
        stift.text(x, y, zeile, size=8, color=MUTED)
        y += 11
    y += 10
    aufwand = a.aufwand_total or 1
    for nr, name, betrag in a.kosten_nach_konto:
        stift.text(x, y, f"{nr} {name}", size=9.5, bold=True)
        stift.text(xr, y, chf(betrag), size=9.5, align="right")
        y += 5
        _balken(stift, x, y, xr - x, betrag / aufwand, GREEN)
        y += 14
        stift.text(x, y, f"{betrag / aufwand * 100:.1f} % vom Aufwand", size=7.5, color=MUTED)
        y += 16

    fusszeile(stift, lager, 4)


def baue_uebersicht(doc: pymupdf.Document, a: Auswertung) -> None:
    zeichne_seite_aufwand(doc, a)
    zeichne_seite_ertrag(doc, a)
    zeichne_seite_pro_person(doc, a)
    zeichne_seite_zahlungen(doc, a)


def baue_pdf(
    a: Auswertung,
    anhang_konten: list[Konto],
    belege_andere: list[tuple[Buchung, Path]],
    belege_lm: list[tuple[Buchung, Path]],
) -> pymupdf.Document:
    doc = pymupdf.open()
    baue_uebersicht(doc, a)

    if anhang_konten:
        trennerseite(
            doc,
            "Kontoauszüge",
            "Nur Konten mit Buchungen dieser Kostenstelle",
        )
        haenge_kontoauszuege_an(doc, anhang_konten)

    if belege_andere or belege_lm:
        trennerseite(
            doc,
            "Belege",
            "Nach Ausgabengrösse — höchste Ausgabe zuerst",
        )
        for b, pfad in belege_andere:
            haenge_datei_an(doc, pfad, _beleg_beschriftung(b))
        if belege_lm:
            trennerseite(doc, "Belege — Lebensmittel", "Konto 4200")
            for b, pfad in belege_lm:
                haenge_datei_an(doc, pfad, _beleg_beschriftung(b))
    return doc


# ─────────────────────────────────────────────────────────────
# CLI
# ─────────────────────────────────────────────────────────────

def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Baut die HeadWaters-Lagerabrechnung als PDF."
    )
    p.add_argument("--kostenstelle", default="C-BTP")
    p.add_argument("--lagername", default="")
    p.add_argument("--ort", default="")
    p.add_argument("--anreise", default="", help="TT.MM.JJJJ, z. B. 22.05.2026")
    p.add_argument("--abreise", default="", help="TT.MM.JJJJ, z. B. 25.05.2026")
    p.add_argument(
        "--teilnehmer",
        type=int,
        default=None,
        help="Erwachsene / Jugendliche, ohne Kinder",
    )
    p.add_argument("--kinder", type=int, default=None)
    p.add_argument(
        "--kind-faktor",
        type=float,
        default=None,
        help="Kinder zählen als so viele Personen (typisch 0.5)",
    )
    p.add_argument("--verpflegungstage", type=float, default=None)
    p.add_argument(
        "--ja",
        action="store_true",
        help="Nicht nachfragen: lagerinfo.json oder alle Flags müssen vollständig sein",
    )
    return p.parse_args()


def _datum_oder_none(text: str) -> date | None:
    if not text:
        return None
    try:
        return parse_datum(text)
    except ValueError:
        raise SystemExit(f"Ungültigs Datum: {text}  (bitte TT.MM.JJJJ)")


def lager_aus_teilen(
    *,
    name: str,
    kostenstelle: str,
    ort: str,
    anreise: date | None,
    abreise: date | None,
    teilnehmer: int | None,
    kinder: int | None,
    kind_faktor: float | None,
    verpflegungstage: float | None,
) -> Lagerinfo | None:
    """Baut nur, wenn wirklich alles da ist — kei stille Standard-Köpfe."""
    if not name or not ort or anreise is None or abreise is None:
        return None
    if teilnehmer is None or kinder is None or verpflegungstage is None:
        return None
    if kind_faktor is None:
        kind_faktor = 0.5
    if abreise <= anreise:
        return None
    if verpflegungstage <= 0:
        return None
    if teilnehmer + kinder <= 0:
        return None
    return Lagerinfo(
        name=name.strip(),
        kostenstelle=kostenstelle.strip(),
        teilnehmer=int(teilnehmer),
        kinder=int(kinder),
        kind_faktor=float(kind_faktor),
        anreise=anreise,
        abreise=abreise,
        verpflegungstage=float(verpflegungstage),
        ort=ort.strip(),
    )


def lagerinfo_abfragen(args: argparse.Namespace, gefundene_ks: str) -> Lagerinfo:
    gespeichert = lagerinfo_laden()
    interaktiv = sys.stdin.isatty() and not args.ja

    print()
    print("  ┌──────────────────────────────────────────┐")
    print("  │   HeadWaters  ·  Lagerabrechnung         │")
    print("  └──────────────────────────────────────────┘")
    print()
    print(f"  Kontoauszüge gefunden für Kostenstelle {gefundene_ks}.")
    print("  Für Esse- und Übernachtigskoste bruuche mer no de Lagersteckbrief.")
    print()

    flag_lager = lager_aus_teilen(
        name=args.lagername or (gespeichert.name if gespeichert and args.ja else ""),
        kostenstelle=args.kostenstelle or gefundene_ks,
        ort=args.ort or (gespeichert.ort if gespeichert and args.ja else ""),
        anreise=_datum_oder_none(args.anreise)
        or (gespeichert.anreise if gespeichert and args.ja else None),
        abreise=_datum_oder_none(args.abreise)
        or (gespeichert.abreise if gespeichert and args.ja else None),
        teilnehmer=args.teilnehmer
        if args.teilnehmer is not None
        else (gespeichert.teilnehmer if gespeichert and args.ja else None),
        kinder=args.kinder
        if args.kinder is not None
        else (gespeichert.kinder if gespeichert and args.ja else None),
        kind_faktor=args.kind_faktor
        if args.kind_faktor is not None
        else (gespeichert.kind_faktor if gespeichert and args.ja else None),
        verpflegungstage=args.verpflegungstage
        if args.verpflegungstage is not None
        else (gespeichert.verpflegungstage if gespeichert and args.ja else None),
    )

    if not interaktiv:
        # --ja: zuerst kompletti Flags, sunsch die gspeichereti Datei
        lager = flag_lager or (gespeichert if args.ja else None)
        if lager is None:
            raise SystemExit(
                "Lagersteckbrief unvollständig.\n"
                "  Entweder uv run generate_abrechnung.py  (interaktiv)\n"
                "  oder alle Flags: --lagername --ort --anreise --abreise "
                "--teilnehmer --kinder --verpflegungstage"
            )
        if flag_lager:
            lagerinfo_speichern(lager)
        print(lager_zusammenfassung(lager))
        print()
        return lager

    vorschlag = gespeichert
    while True:
        print("  ── Lagersteckbrief ─────────────────────────")
        print("     Enter = Vorschlag, wo vorhanden.")
        print()

        name = frage(
            "  Lagername",
            args.lagername or (vorschlag.name if vorschlag else "Bibeltage Pfingsten 2026"),
            pflicht=True,
        )
        kostenstelle = frage(
            "  Kostenstelle",
            args.kostenstelle or gefundene_ks,
            pflicht=True,
        )
        ort = frage(
            "  Ort / Unterkunft",
            args.ort or (vorschlag.ort if vorschlag else ""),
            pflicht=True,
        )
        anreise = frage_datum(
            "  Anreise (TT.MM.JJJJ)",
            _datum_oder_none(args.anreise)
            or (vorschlag.anreise if vorschlag else None),
        )
        abreise = frage_datum(
            "  Abreise (TT.MM.JJJJ)",
            _datum_oder_none(args.abreise)
            or (vorschlag.abreise if vorschlag else None),
        )
        if abreise <= anreise:
            print("   → D Abreise mues nach der Anreise si.")
            continue
        naechte = (abreise - anreise).days
        print(f"     → {naechte} Nächte, {format_spanne(anreise, abreise)}")
        print()
        print("  Verpflegungstage: An- und Abreisetag zelle oft nume ½.")
        print("  Bispil: 3 Nächt, Frühstück Sa bis Mittag Mo → 2.5")
        vorschlag_tage = (
            args.verpflegungstage
            if args.verpflegungstage is not None
            else (
                vorschlag.verpflegungstage
                if vorschlag
                else max(0.5, naechte - 0.5)
            )
        )
        verpflegungstage = frage_zahl(
            "  Verpflegungstage",
            vorschlag_tage,
        )
        print()
        teilnehmer = frage_zahl(
            "  Anzahl Erwachsene / Jugendliche (ohne Kinder)",
            args.teilnehmer
            if args.teilnehmer is not None
            else (vorschlag.teilnehmer if vorschlag else None),
            ganzzahl=True,
        )
        kinder = frage_zahl(
            "  Anzahl Kinder",
            args.kinder
            if args.kinder is not None
            else (vorschlag.kinder if vorschlag else 0),
            ganzzahl=True,
        )
        kind_faktor = frage_zahl(
            "  Kinder zelle als wieviil Personen?",
            args.kind_faktor
            if args.kind_faktor is not None
            else (vorschlag.kind_faktor if vorschlag else 0.5),
        )

        lager = lager_aus_teilen(
            name=name,
            kostenstelle=kostenstelle,
            ort=ort,
            anreise=anreise,
            abreise=abreise,
            teilnehmer=teilnehmer,
            kinder=kinder,
            kind_faktor=kind_faktor,
            verpflegungstage=verpflegungstage,
        )
        if lager is None:
            print("   → Da fählt öppis oder d Zahle stimme nid. No einisch.")
            continue

        print()
        print("  ── Passt das? ──────────────────────────────")
        print(lager_zusammenfassung(lager))
        print()
        ok = frage("  Stimmt das? J/n", "J")
        if ok.lower() in {"n", "nein", "no"}:
            vorschlag = lager
            print()
            continue

        lagerinfo_speichern(lager)
        print(f"  Gspeicheret under {LAGERINFO_DATEI.name} (nächschte Lauf als Vorschlag).")
        print()
        return lager


def main() -> None:
    args = parse_args()

    if not DATA.is_dir():
        raise SystemExit(f"Data-Ordner fehlt: {DATA}")

    xlsx_dir = next(
        (
            p
            for p in DATA.iterdir()
            if p.is_dir()
            and "2026" in p.name
            and "2026-2" not in p.name
            and "2026-3" not in p.name
        ),
        None,
    )
    pdf_dir = next(
        (p for p in DATA.iterdir() if p.is_dir() and "2026-2" in p.name),
        None,
    )
    eingang_dir = next(
        (p for p in DATA.iterdir() if p.is_dir() and "2026-3" in p.name),
        None,
    )
    if not xlsx_dir or not (xlsx_dir / "Buchungsjournal.xlsx").exists():
        raise SystemExit("Kei Excel-Ordner «Kontoauszüge Buchungsperiode 2026» gfunde.")
    if not pdf_dir:
        raise SystemExit("Kei PDF-Ordner «Kontoauszüge … 2026-2» gfunde.")

    print(f"Excel (Kostenstelle): {xlsx_dir.name}")
    print(f"PDFs:                 {pdf_dir.name}")
    if eingang_dir:
        print(f"Excel (Eingänge):     {eingang_dir.name}")

    lager = lagerinfo_abfragen(args, args.kostenstelle or "C-BTP")
    print()
    print(f"→ Lese Buchungen für {lager.kostenstelle} …")

    buchungen = lies_journal(xlsx_dir, lager.kostenstelle)
    if not buchungen:
        raise SystemExit(f"Kei Buchige mit Kostenstelle {lager.kostenstelle} im Journal.")
    print(f"  {len(buchungen)} Buchungen im Journal")

    eingaenge: list[Buchung] = []
    rechnungen: set[str] = set()
    if eingang_dir and (eingang_dir / "Buchungsjournal.xlsx").exists():
        eingaenge = lies_journal(eingang_dir, kostenstelle=None)
        rechnungen = rechnungsnummern_aus_xlsx(xlsx_dir, lager.kostenstelle)
        n_codes = len(paycodes_aus_buchungen(buchungen))
        print(
            f"  {len(eingaenge)} Eingänge ohne Kostenstelle · "
            f"{n_codes} Paycodes · {len(rechnungen)} Rechnungsnummern"
        )

    konten = sammle_konten(xlsx_dir, pdf_dir, lager.kostenstelle)
    auswertung = auswerten(lager, buchungen, konten, eingaenge, rechnungen)

    anhang = pdfs_zum_anhaengen(konten)
    print(f"  {len(anhang)} Kontoauszug-PDFs mit Buchungen (leere fliegen raus)")

    beleg_index = index_belege(BELEGE_DIR)
    belege_andere, belege_lm = belege_zum_anhaengen(buchungen, beleg_index)
    print(
        f"  {len(belege_andere) + len(belege_lm)} Belege dieser Kostenstelle"
        f" ({len(belege_lm)} Lebensmittel)"
    )

    if not LOGO_DATEI.exists():
        raise SystemExit(f"Logo fehlt: {LOGO_DATEI}")

    OUTPUT.mkdir(parents=True, exist_ok=True)
    print("→ Baue PDF …")
    doc = baue_pdf(auswertung, anhang, belege_andere, belege_lm)
    slug = re.sub(r"[^A-Za-z0-9]+", "_", lager.name).strip("_")
    ziel = OUTPUT / f"Abrechnung_{slug}.pdf"
    doc.save(ziel, garbage=4, deflate=True)
    n_seiten = doc.page_count
    doc.close()

    print()
    print("  Fertig.")
    print(f"  {ziel}  ({n_seiten} Seiten)")
    print()
    print("  Aufwand:          ", chf(auswertung.aufwand_total))
    print("  Ertrag:           ", chf(auswertung.ertrag_total))
    print("  Saldo dieses Jahr:", chf(auswertung.saldo_jahr))
    print("  Übriges Geld:     ", chf(auswertung.uebriges_geld))
    print("  Zahlungen:")
    for m in auswertung.zahlungen:
        print(
            f"    {m.name:28} {m.anzahl:3}×  {chf(m.volumen):>14}"
            f"   Gebühr {chf(m.gebuehren)}"
        )


if __name__ == "__main__":
    main()
